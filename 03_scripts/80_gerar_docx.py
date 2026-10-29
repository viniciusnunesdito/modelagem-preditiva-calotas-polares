# -*- coding: utf-8 -*-
"""
Etapa 80: geracao do artigo em formato .docx conforme o modelo institucional.

Formatacao aplicada (Modelo Institucional de Artigo Cientifico, Prof. Dr. Renato
Goncalves Rocha): papel A4; margens superior e esquerda 3 cm, inferior e direita 2 cm;
fonte Arial 12 no corpo; espacamento entre linhas 1,5; texto justificado com recuo de
primeira linha de 1,25 cm; titulos de secao em Arial 12 negrito alinhados a esquerda,
com numeracao progressiva; figuras e tabelas numeradas, com titulo e fonte.
"""
import os, re, json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ART = "06_artigo/ARTIGO.md"
FIG = "04_figuras"
SAIDA = "06_artigo/ARTIGO_Modelagem_Preditiva_Calotas_Polares.docx"

FIGURAS_COMPLETO = {
    1: ("Figura_01_fluxograma_metodologico.png",
        "Estrutura metodológica da pesquisa e seus cinco pontos de validação.", 15.2),
    2: ("Figura_02_impressoes_digitais.png",
        "Fatores de impressão digital do nível do mar: (a) transecto ao longo do litoral "
        "brasileiro, de 4,4° N a 34° S; (b) contraste entre campo próximo e campo distante.", 16.0),
    3: ("Figura_03_fechamento_orcamento.png",
        "Fechamento do orçamento do nível do mar: (a) escala global e de bacia; "
        "(b) estações brasileiras, com intervalo de confiança de 95% com erro padrão robusto.", 16.0),
    4: ("Figura_04_benchmark_algoritmos.png",
        "Comparação de quinze algoritmos preditivos sob validação de origem móvel: (a) erro "
        "quadrático médio fora da amostra no horizonte de 20 anos, nas duas estações brasileiras "
        "de registro longo, com as linhas verticais marcando a previsão trivial pela média de cada "
        "uma e as setas indicando barras truncadas em 250 mm; (b) comportamento na extrapolação ao "
        "ano 2100, em que a linha cinza marca a contribuição observada em 2018.", 16.0),
    5: ("Figura_05_validacao_emulador.png",
        "Validação V3: emulador calibrado apenas com 1900–1989 e solicitado a prever 1990–2018, "
        "comparado à extrapolação linear ajustada na mesma janela.", 16.0),
    6: ("Figura_06_sensibilidade_projecoes.png",
        "(a) Sensibilidade das calotas ao aquecimento medida em três épocas; "
        "(b) projeção da elevação média global em 2100 nas três vias.", 16.0),
    7: ("Figura_07_projecoes_locais.png",
        "(a) Decomposição da elevação projetada do nível relativo do mar em 2100 por estação; "
        "(b) parcela atribuível às calotas polares, por cenário.", 16.0),
    8: ("Figura_08_extremos_recuo.png",
        "(a) Curvas de período de retorno na Ilha Fiscal deslocadas pela elevação projetada; "
        "(b) recuo potencial da linha de costa pela regra de Bruun.", 16.0),
}

# mapeamento da versao condensada: numero no artigo -> figura gerada
FIGURAS_ARTIGO = {
    1: FIGURAS_COMPLETO[1],
    2: FIGURAS_COMPLETO[2],
    3: FIGURAS_COMPLETO[3],
    4: FIGURAS_COMPLETO[4],
    5: FIGURAS_COMPLETO[7],
    6: FIGURAS_COMPLETO[8],
}

import sys as _sys
if len(_sys.argv) > 1 and _sys.argv[1] == 'completo':
    ART = "06_artigo/RELATORIO_CIENTIFICO_COMPLETO.md"
    SAIDA = "06_artigo/RELATORIO_CIENTIFICO_COMPLETO.docx"
    FIGURAS = FIGURAS_COMPLETO
else:
    ART = "06_artigo/ARTIGO.md"
    SAIDA = "06_artigo/ARTIGO_Modelagem_Preditiva_Calotas_Polares.docx"
    FIGURAS = FIGURAS_ARTIGO


def arial(run, tam=12, negrito=False, italico=False, cor=None):
    run.font.name = 'Arial'
    run.font.size = Pt(tam)
    run.bold = negrito
    run.italic = italico
    if cor:
        run.font.color.rgb = cor
    r = run._element.rPr.rFonts
    r.set(qn('w:eastAsia'), 'Arial')
    r.set(qn('w:cs'), 'Arial')


def par(doc, texto='', tam=12, negrito=False, italico=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        recuo=0.0, antes=0, depois=6, espaco=1.5, cor=None):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.first_line_indent = Cm(recuo)
    pf.space_before = Pt(antes)
    pf.space_after = Pt(depois)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE if espaco == 1.5 else WD_LINE_SPACING.SINGLE
    if espaco not in (1.0, 1.5):
        pf.line_spacing = espaco
    if texto:
        escreve_formatado(p, texto, tam, negrito, italico, cor)
    return p


def escreve_formatado(p, texto, tam=12, negrito=False, italico=False, cor=None):
    """Interpreta **negrito** e *italico* do markdown."""
    partes = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', texto)
    for t in partes:
        if not t:
            continue
        if t.startswith('**') and t.endswith('**'):
            arial(p.add_run(t[2:-2]), tam, True, italico, cor)
        elif t.startswith('*') and t.endswith('*') and len(t) > 2:
            arial(p.add_run(t[1:-1]), tam, negrito, True, cor)
        else:
            arial(p.add_run(t), tam, negrito, italico, cor)


def sombreia(celula, cor_hex):
    tcPr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), cor_hex)
    tcPr.append(shd)


# ------------------------------------------------------------------ documento
doc = Document()
st = doc.styles['Normal']
st.font.name = 'Arial'; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin, sec.left_margin = Cm(3.0), Cm(3.0)
sec.bottom_margin, sec.right_margin = Cm(2.0), Cm(2.0)

texto = open(ART, encoding='utf-8').read()
linhas = texto.split('\n')

i = 0
n_fig = 0
n_tab = 0
em_tabela = False
buffer_tab = []


def descarrega_tabela(doc, linhas_tab):
    """Converte um bloco markdown de tabela em tabela do Word."""
    global n_tab
    linhas_tab = [l for l in linhas_tab if l.strip().startswith('|')]
    if len(linhas_tab) < 3:
        return
    corpo = [l for l in linhas_tab if not re.match(r'^\s*\|[\s:|-]+\|\s*$', l)]
    celulas = [[c.strip() for c in l.strip().strip('|').split('|')] for l in corpo]
    ncol = max(len(r) for r in celulas)
    celulas = [r + [''] * (ncol - len(r)) for r in celulas]
    t = doc.add_table(rows=len(celulas), cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, linha in enumerate(celulas):
        for ci, val in enumerate(linha):
            cel = t.cell(ri, ci)
            cel.text = ''
            p = cel.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (ci > 0 or ri == 0) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            escreve_formatado(p, val, tam=9, negrito=(ri == 0))
            if ri == 0:
                sombreia(cel, 'EFEFEC')


while i < len(linhas):
    l = linhas[i]
    s = l.strip()

    # --------- tabelas markdown
    if s.startswith('|'):
        buffer_tab.append(l)
        i += 1
        continue
    elif buffer_tab:
        descarrega_tabela(doc, buffer_tab)
        buffer_tab = []

    # --------- titulo principal
    if s.startswith('# '):
        p = par(doc, s[2:].strip(), tam=14, negrito=True,
                align=WD_ALIGN_PARAGRAPH.CENTER, antes=0, depois=18)
    # --------- secoes
    elif s.startswith('## '):
        t = s[3:].strip()
        par(doc, t, tam=12, negrito=True, align=WD_ALIGN_PARAGRAPH.LEFT,
            antes=18, depois=10)
    elif s.startswith('### '):
        par(doc, s[4:].strip(), tam=12, negrito=True, align=WD_ALIGN_PARAGRAPH.LEFT,
            antes=12, depois=8)
    # --------- separador
    elif s == '---':
        pass
    # --------- marcador de figura
    elif s.startswith('**[Figura'):
        n_fig += 1
        if n_fig in FIGURAS:
            arq, legenda, larg = FIGURAS[n_fig]
            cam = os.path.join(FIG, arq)
            if os.path.exists(cam):
                pimg = doc.add_paragraph()
                pimg.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pimg.paragraph_format.space_before = Pt(10)
                pimg.paragraph_format.space_after = Pt(4)
                pimg.add_run().add_picture(cam, width=Cm(larg))
                p = par(doc, '', tam=10, align=WD_ALIGN_PARAGRAPH.CENTER,
                        antes=0, depois=2, espaco=1.0)
                arial(p.add_run(f'Figura {n_fig} – '), 10, True)
                arial(p.add_run(legenda), 10)
        # pula a linha "Fonte:" seguinte, tratada abaixo
    elif s.startswith('Fonte: Elaborado pelos autores'):
        p = par(doc, '', tam=10, align=WD_ALIGN_PARAGRAPH.CENTER,
                antes=0, depois=12, espaco=1.0)
        arial(p.add_run(s), 10)
    elif s.startswith('**Tabela'):
        n_tab += 1
        m = re.match(r'\*\*Tabela\s+(\d+)\s*[–-]\s*(.+?)\*\*', s)
        titulo = m.group(2) if m else s.strip('*')
        p = par(doc, '', tam=10, align=WD_ALIGN_PARAGRAPH.LEFT,
                antes=12, depois=3, espaco=1.0)
        arial(p.add_run(f'Tabela {n_tab} – '), 10, True)
        arial(p.add_run(titulo), 10)
    # --------- texto comum
    elif s:
        prim = s.split(' ')[0]
        # paragrafos de metadados do cabecalho (autores, afiliacoes)
        if s.startswith(('Nome Completo', 'Prof. Dr.', '¹ Curso', '² Professor',
                         'Autor correspondente')):
            par(doc, s, tam=11, align=WD_ALIGN_PARAGRAPH.CENTER, antes=0, depois=3, espaco=1.0)
        elif s.startswith('**Palavras-chave') or s.startswith('**Keywords'):
            par(doc, s, tam=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, antes=4, depois=10)
        elif prim in ('RESUMO', 'ABSTRACT'):
            par(doc, s, tam=12, negrito=True, align=WD_ALIGN_PARAGRAPH.LEFT, antes=14, depois=8)
        else:
            # resumo e abstract em paragrafo unico sem recuo
            recuo = 0.0 if len(s) > 900 else 1.25
            par(doc, s, tam=12, recuo=recuo)
    i += 1

if buffer_tab:
    descarrega_tabela(doc, buffer_tab)

# --- metadados do documento -------------------------------------------------
# A biblioteca usada para montar o arquivo grava, por padrao, 'python-docx' no
# campo de autor e uma data de 2013 herdada do seu modelo interno. Esses campos
# ficam visiveis nas propriedades do arquivo e sao sobrescritos aqui. O campo de
# autor fica em branco para ser preenchido pelos autores no editor de texto.
import datetime as _dt
cp = doc.core_properties
cp.author = ''
cp.last_modified_by = ''
cp.comments = ''
cp.category = ''
cp.keywords = ('elevacao do nivel do mar; impressao digital do nivel do mar; '
               'modelagem preditiva; erosao costeira; validacao de modelos')
cp.title = ('Modelagem preditiva dos impactos do derretimento das calotas polares '
            'sobre a dinamica costeira brasileira')
cp.subject = 'Artigo cientifico - Tendencias da Computacao'
cp.created = _dt.datetime.now()
cp.modified = _dt.datetime.now()
cp.revision = 1

doc.save(SAIDA)
print(f"gravado {SAIDA}")
print(f"  {n_fig} figuras inseridas, {n_tab} tabelas numeradas")
print(f"  tamanho: {os.path.getsize(SAIDA)/1e6:.2f} MB")
